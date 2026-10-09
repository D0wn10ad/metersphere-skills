#!/usr/bin/env python3
"""测试用例文件拆分器：docx/pdf/xlsx/xmind → v2 字段契约 JSON 数组草稿。

stdlib-only 核心 + 可选富依赖优雅降级（模式同 ms_import_helper.py）：
- xmind：stdlib zipfile+json/xml 解析（无富依赖），按 v2.10 XmindCaseParser 语义。
- xlsx：需 openpyxl（缺省报错 + 安装提示）。
- docx：需 python-docx（缺省报错 + 安装提示）。
- pdf：需 pdfplumber（缺省报错 + 安装提示）。

输出草稿供 AI 增强（references/ai-phabricator-functional-case-prompt.md）
或直接 functional-case batch-create。每条草稿必含 name/nodePath/priority/steps，
缺关键列/关键节点的行跳过并在 stderr 告警计数。

字段契约（v2，禁 v3-only 字段 moduleId/customNum/maintainer/customFields/aiCreate）：
必填 name/nodePath/priority(P0-P3)/steps([{num,desc,result}])/caseEditType(TEXT|STEP)；
可选仅在有值时输出 description/precondition/remark/tags（JSON 数组字符串）。
xmind 的 id: 子节点仅 Update/useCustomId 生效（v2.10 XmindCaseParser），Create 草稿省略。
不调用网络；报错全 zh-CN。
"""
import argparse
import json
import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree

DEFAULT_NODE_PATH = "/默认模块"
RICH_DEPS = {"xlsx": ("openpyxl", "openpyxl"),
             "docx": ("docx", "python-docx"),
             "pdf": ("pdfplumber", "pdfplumber")}
SUPPORTED_FORMATS = ("xlsx", "docx", "pdf", "xmind")
XLSX_HEADER_MAP = {"用例名称": "name", "所属模块": "nodePath", "标签": "tags",
                   "前置条件": "precondition", "步骤描述": "step_desc",
                   "预期结果": "step_result", "编辑模式": "caseEditType",
                   "备注": "remark", "用例等级": "priority"}
TC_CHILD_RE = re.compile(r"^(pc|rc|tag|id)\s*[:：]\s*(.*)$", re.I)


def _require_dep(fmt):
    module_name, pip_name = RICH_DEPS[fmt]
    try:
        return __import__(module_name)
    except ImportError:
        raise ValueError(
            f"解析 {fmt} 文件需要 {pip_name}：请先安装（pip install {pip_name}）后重试")


def detect_format(path):
    suffix = Path(str(path)).suffix.lower()
    for fmt in SUPPORTED_FORMATS:
        if suffix == "." + fmt:
            return fmt
    raise ValueError(
        f"无法按扩展名识别文件格式: {suffix or '(无)'}（支持 {', '.join(SUPPORTED_FORMATS)}），"
        f"请用 --format 显式指定")


def parse_tc_title(title):
    """tc 前缀解析（镜像 v2.10 XmindCaseParser：前缀 strip 后为用例名）。

    `tc:登录`/`tc：登录`/`tc登录` → ("P0", "登录")；`tc-P1:登录` → ("P1", "登录")。
    非 tc 前缀或剩余名为空 → None。
    """
    t = (title or "").strip()
    if len(t) < 2 or t[:2].lower() != "tc":
        return None
    rest = t[2:]
    priority = "P0"
    m = re.match(r"^[-_ ]?(P[0-3])(?![0-9])\s*[:：]?\s*", rest, re.I)
    if m:
        priority = m.group(1).upper()
        rest = rest[m.end():]
    else:
        rest = re.sub(r"^\s*[:：]\s*", "", rest)
    rest = rest.strip()
    if not rest:
        return None
    return priority, rest


def _normalize_priority(value):
    p = str(value or "").strip().upper()
    return p if re.fullmatch(r"P[0-3]", p) else "P0"


def _normalize_tags(value):
    """标签列：JSON 数组字符串（服务端契约）或普通文本（逗号/中文逗号分隔）。"""
    raw = str(value or "").strip()
    if not raw:
        return None
    try:
        parsed = json.loads(raw)
        if isinstance(parsed, list):
            tags = [str(t).strip() for t in parsed if str(t).strip()]
            return json.dumps(tags, ensure_ascii=False) if tags else None
    except json.JSONDecodeError:
        pass
    tags = [t.strip() for t in re.split(r"[,，]", raw) if t.strip()]
    return json.dumps(tags, ensure_ascii=False) if tags else None


def _finalize(drafts):
    valid, skipped = [], 0
    for d in drafts:
        missing = [k for k in ("name", "nodePath", "priority", "steps")
                   if not d.get(k)]
        if missing:
            skipped += 1
            print(f"警告: 跳过 1 条草稿（缺 {', '.join(missing)}）", file=sys.stderr)
            continue
        valid.append(d)
    if skipped:
        print(f"警告: 共跳过 {skipped} 条草稿（缺关键列/关键节点）", file=sys.stderr)
    return valid


def _draft(name, node_path, priority, steps, case_edit_type="STEP",
           precondition=None, remark=None, tags=None, project_id=None):
    d = {"name": str(name).strip(), "nodePath": str(node_path).strip(),
         "priority": _normalize_priority(priority),
         "steps": steps, "caseEditType": case_edit_type}
    if precondition:
        d["precondition"] = str(precondition).strip()
    if remark:
        d["remark"] = str(remark).strip()
    if tags:
        d["tags"] = tags
    if project_id:
        d["projectId"] = str(project_id).strip()
    return d


def _split_xlsx(path, project_id=None):
    _require_dep("xlsx")
    from openpyxl import load_workbook
    wb = load_workbook(str(path), read_only=True, data_only=True)
    ws = wb.active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    if not rows:
        raise ValueError("xlsx 文件无内容行")
    header_idx = {}
    for i, cell in enumerate(rows[0]):
        name = str(cell or "").strip()
        if name in XLSX_HEADER_MAP and XLSX_HEADER_MAP[name] not in header_idx:
            header_idx[XLSX_HEADER_MAP[name]] = i
    if "name" not in header_idx:
        raise ValueError("xlsx 首行未找到「用例名称」列表头（须为 MeterSphere 导入模板格式）")
    col = lambda row, key: row[header_idx[key]] if header_idx.get(key) is not None and header_idx[key] < len(row) else None
    drafts, current = [], None
    for row in rows[1:]:
        name = str(col(row, "name") or "").strip()
        step_desc = str(col(row, "step_desc") or "").strip()
        has_content = any(str(c or "").strip() for c in row)
        if not name and current is not None and step_desc:
            current["steps"].append({"num": len(current["steps"]) + 1,
                                     "desc": step_desc,
                                     "result": str(col(row, "step_result") or "").strip()})
            continue
        if current is not None:
            drafts.append(current)
            current = None
        if not name:
            if has_content:
                drafts.append(_draft("", DEFAULT_NODE_PATH, "P0", [],
                                     project_id=project_id))
            continue
        node_path = str(col(row, "nodePath") or "").strip() or DEFAULT_NODE_PATH
        edit_type = str(col(row, "caseEditType") or "").strip().upper()
        current = _draft(name, node_path, col(row, "priority"),
                         [{"num": 1, "desc": step_desc,
                           "result": str(col(row, "step_result") or "").strip()}] if step_desc else [],
                         edit_type if edit_type in ("TEXT", "STEP") else "STEP",
                         precondition=col(row, "precondition"),
                         remark=col(row, "remark"),
                         tags=_normalize_tags(col(row, "tags")),
                         project_id=project_id)
    if current is not None:
        drafts.append(current)
    return _finalize(drafts)


def _xmind_nodes(path):
    """xmind → 根节点树 [{title, children}]（content.json 新版 / content.xml 旧版）。"""
    with zipfile.ZipFile(str(path)) as zf:
        names = zf.namelist()
        if "content.json" in names:
            doc = json.loads(zf.read("content.json").decode("utf-8"))
            sheets = doc if isinstance(doc, list) else [doc]

            def from_json(topic):
                return {"title": str(topic.get("title") or ""),
                        "children": [from_json(c) for c in
                                     ((topic.get("children") or {}).get("attached") or [])]}
            return [from_json(s.get("rootTopic") or {}) for s in sheets
                    if isinstance(s, dict)]
        if "content.xml" in names:
            root = ElementTree.fromstring(zf.read("content.xml"))

            def from_xml(topic_el):
                return {"title": str(topic_el.findtext("title") or ""),
                        "children": [from_xml(t) for t in
                                     topic_el.findall("./children/topics/topic")]}
            return [from_xml(sheet.find("./topic"))
                    for sheet in root.findall("sheet")
                    if sheet.find("./topic") is not None]
    raise ValueError("xmind 文件缺少 content.json/content.xml（非合法 xmind 包）")


def _split_xmind(path, project_id=None):
    drafts = []

    def walk(node, module_path):
        tc = parse_tc_title(node["title"])
        if tc is None and node["title"].strip()[:2].lower() == "tc":
            print("警告: 跳过 1 条 xmind 用例节点（tc 前缀但用例名为空）", file=sys.stderr)
        if tc is not None:
            priority, name = tc
            steps, precondition, remark, tags = [], None, None, None
            for child in node["children"]:
                m = TC_CHILD_RE.match(child["title"].strip())
                if m:
                    kind, value = m.group(1).lower(), m.group(2).strip()
                    if kind == "pc":
                        precondition = value
                    elif kind == "rc":
                        remark = value
                    elif kind == "tag":
                        tags = _normalize_tags(value)
                elif child["title"].strip():
                    steps.append({"num": len(steps) + 1,
                                  "desc": child["title"].strip(),
                                  "result": (child["children"][0]["title"].strip()
                                             if child["children"] else "")})
            drafts.append(_draft(name, module_path or DEFAULT_NODE_PATH,
                                 priority, steps, precondition=precondition,
                                 remark=remark, tags=tags, project_id=project_id))
            return
        child_path = module_path
        if node["title"].strip():
            child_path = module_path + "/" + node["title"].strip()
        for child in node["children"]:
            walk(child, child_path)

    for root in _xmind_nodes(path):
        walk(root, "")
    return _finalize(drafts)


def _iter_docx_blocks(doc):
    """按文档顺序产出段落与表格（python-docx 无序迭代问题的版本安全解法）。"""
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, doc)
        elif child.tag.endswith("}tbl"):
            yield Table(child, doc)


def _split_docx(path, project_id=None):
    _require_dep("docx")
    import docx
    doc = docx.Document(str(path))
    drafts, current = [], None
    module_path = ""
    for block in _iter_docx_blocks(doc):
        if hasattr(block, "rows"):
            for row in block.rows:
                cells = [c.text.strip() for c in row.cells]
                if cells and cells[0] in ("步骤描述", "desc", "用例名称"):
                    continue
                if current is not None and any(cells):
                    current["steps"].append({
                        "num": len(current["steps"]) + 1,
                        "desc": cells[0] if cells else "",
                        "result": cells[1] if len(cells) > 1 else ""})
            continue
        style = (block.style.name or "") if block.style is not None else ""
        text = block.text.strip()
        m = re.match(r"^(?:Heading|标题)\s*([1-9])$", style, re.I)
        if m and text:
            level = int(m.group(1))
            if level <= 2:
                module_path = ("/" + text) if level == 1 else (module_path + "/" + text)
                continue
            if current is not None:
                drafts.append(current)
            tc = parse_tc_title(text)
            if tc is not None:
                priority, name = tc
                current = _draft(name, module_path or DEFAULT_NODE_PATH, priority,
                                 [], project_id=project_id)
            else:
                current = _draft(text, module_path or DEFAULT_NODE_PATH, "P0", [],
                                 project_id=project_id)
            continue
        tc = parse_tc_title(text) if text else None
        if tc is not None:
            if current is not None:
                drafts.append(current)
            priority, name = tc
            current = _draft(name, module_path or DEFAULT_NODE_PATH, priority,
                             [], project_id=project_id)
            continue
        if not text:
            continue
        pm = re.match(r"^(前置条件|备注)\s*[:：]\s*(.*)$", text)
        if pm and current is not None:
            if pm.group(1) == "前置条件":
                current["precondition"] = pm.group(2).strip()
            else:
                current["remark"] = pm.group(2).strip()
        elif current is not None:
            current["steps"].append({"num": len(current["steps"]) + 1,
                                     "desc": text, "result": ""})
    if current is not None:
        drafts.append(current)
    return _finalize(drafts)


def _split_pdf(path, project_id=None):
    _require_dep("pdf")
    import pdfplumber
    drafts, current = [], None
    module_path = ""
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            for line in (page.extract_text() or "").splitlines():
                text = line.strip()
                if not text:
                    continue
                hm = re.match(r"^(#{1,2})\s+(.+)$", text)
                if hm:
                    module_path = ("/" + hm.group(2).strip()) if len(hm.group(1)) == 1 \
                        else (module_path + "/" + hm.group(2).strip())
                    continue
                tc = parse_tc_title(text)
                if tc is not None:
                    if current is not None:
                        drafts.append(current)
                    priority, name = tc
                    current = _draft(name, module_path or DEFAULT_NODE_PATH,
                                     priority, [], project_id=project_id)
                    continue
                pm = re.match(r"^(前置条件|备注)\s*[:：]\s*(.*)$", text)
                if pm and current is not None:
                    if pm.group(1) == "前置条件":
                        current["precondition"] = pm.group(2).strip()
                    else:
                        current["remark"] = pm.group(2).strip()
                elif current is not None:
                    current["steps"].append({"num": len(current["steps"]) + 1,
                                             "desc": text, "result": ""})
    if current is not None:
        drafts.append(current)
    return _finalize(drafts)


SPLITTERS = {"xlsx": _split_xlsx, "docx": _split_docx,
             "pdf": _split_pdf, "xmind": _split_xmind}


def split_file(path, fmt="auto", project_id=None):
    source = Path(str(path)).expanduser()
    if not source.exists():
        raise ValueError(f"文件不存在: {path}")
    if fmt == "auto":
        fmt = detect_format(path)
    elif fmt not in SPLITTERS:
        raise ValueError(f"不支持的格式: {fmt}（支持 {', '.join(SUPPORTED_FORMATS)}）")
    return SPLITTERS[fmt](str(source), project_id=project_id)


def main(argv=None):
    parser = argparse.ArgumentParser(
        description="测试用例文件拆分器：docx/pdf/xlsx/xmind → v2 字段契约 JSON 草稿")
    parser.add_argument("file", help="测试用例文件（docx/pdf/xlsx/xmind）")
    parser.add_argument("--format", default="auto",
                        choices=["auto"] + list(SUPPORTED_FORMATS),
                        help="文件格式（默认 auto 按扩展名推断）")
    parser.add_argument("-o", "--out", default=None, help="输出 JSON 文件（默认 stdout）")
    parser.add_argument("--project-id", default=None,
                        help="可选：注入每条草稿的 projectId（缺省由 AI 增强或调用方填充）")
    args = parser.parse_args(argv)
    try:
        drafts = split_file(args.file, args.format, project_id=args.project_id)
    except ValueError as e:
        print(f"错误: {e}", file=sys.stderr)
        return 1
    payload = json.dumps(drafts, ensure_ascii=False, indent=2)
    if args.out:
        with open(args.out, "w", encoding="utf-8") as fh:
            fh.write(payload + "\n")
        print(f"已拆分 {len(drafts)} 条草稿到 {args.out}")
    else:
        print(payload)
    return 0


if __name__ == "__main__":
    sys.exit(main())
